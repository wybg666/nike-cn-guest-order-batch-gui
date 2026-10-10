import json
import queue
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import requests
import tkinter as tk
from tkinter import messagebox, ttk
from openpyxl import Workbook

from guest_batch_service import (
    ProxyItem,
    build_failed_row,
    ensure_deferred_form_fields,
    normalize_items,
    normalize_order_row,
    parse_order_phone_lines,
    parse_proxies,
    resolve_payment_approval_id,
)


DEFAULT_ORDER_PHONE_TEXT = ""  # Public release: no personal data or credentials.


DEFAULT_PROXY_TEXT = ""  # Public release: no personal data or credentials.


API_TEMPLATE = "https://api.nike.com.cn/orders/summary/v1/{order_id}"
DEFERRED_PAYMENT_POST_URL = "https://api.nike.com.cn/payment/deferred_payment_forms/v2"
DEFERRED_PAYMENT_JOB_URL = "https://api.nike.com.cn/payment/deferred_payment_forms/v2/jobs/{job_id}"
COMMON_PARAMS = {
    "locale": "zh_cn",
    "country": "CN",
    "language": "zh-Hans",
    "timezone": "Asia/Shanghai",
}
DEFAULT_HEADERS = {
    "accept": "application/json, text/plain, */*",
    "origin": "https://www.nike.com.cn",
    "referer": "https://www.nike.com.cn/orders/details/",
    "nike-api-caller-id": "com.nike.commerce.nikedotcom.web",
    "user-agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
}
SUMMARY_HTTP_OK = {200, 400, 401, 403, 404, 429}
SETTINGS_FILENAME = "nike_guest_batch_gui_settings.json"
DEFAULT_THREAD_COUNT = 5
DEFAULT_MAX_PROXY_ATTEMPTS = 3
DEFAULT_PROXY_PRECHECK = True

ORDER_COLUMNS = [
    "orderId",
    "queryPhone",
    "paymentApprovalId",
    "paymentHoldFlag",
    "status",
    "subStatus",
    "orderDate",
    "currency",
    "orderTotal",
    "chargedPrice",
    "paymentStatus",
    "receiverName",
    "receiverPhone",
    "receiverRegion",
    "receiverAddressLine",
    "receiverAddressLine2",
    "receiverAddressLine3",
    "fullShippingAddress",
    "shippingAddressRaw",
    "deferredFormStatus",
    "deferredFormJobId",
    "deferredFormValidUntil",
    "deferredFormAction",
    "deferredQrCodeURL",
    "deferredFormError",
    "rawJson",
]
ITEM_COLUMNS = [
    "orderId",
    "queryPhone",
    "groupHeading",
    "groupSubheading",
    "groupEncodedMessage",
    "productTitle",
    "productId",
    "productPid",
    "productCategory",
    "productImage",
    "styleNumber",
    "color",
    "styleColor",
    "gtin",
    "size",
    "lineStatus",
    "lineSubStatus",
    "lineStatusColor",
    "lineSubStatusColor",
    "quantity",
    "returnableQuantity",
    "lineItemChargedPrice",
    "lineItemFees",
    "lineItemOrderDate",
    "estimatedDelivery",
    "estimatedDeliveryMaxDateTime",
    "estimatedDeliveryPrecision",
    "timezoneId",
    "orderLineType",
    "linePaymentStatus",
    "sizeDescription",
    "standardCarrierAlphaCode",
    "returnableFlag",
    "groupActionsRaw",
    "itemActionsRaw",
]
FAILED_COLUMNS = ["orderId", "queryPhone", "error", "proxy", "attempt"]


def get_runtime_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


class NikeGuestBatchGUI:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Nike 访客查单批量工具")
        self.root.geometry("1300x860")

        self.log_queue: "queue.Queue[tuple]" = queue.Queue()
        self.stop_event = threading.Event()
        self.worker: Optional[threading.Thread] = None

        self.results_orders: List[dict] = []
        self.results_items: List[dict] = []
        self.failed_rows: List[dict] = []
        self.results_lock = threading.Lock()

        self.total = 0
        self.done = 0
        self.success = 0
        self.last_export_path: Optional[Path] = None
        self.proxy_pool: "queue.Queue[ProxyItem]" = queue.Queue()
        self.proxy_pool_size = 0
        self.settings_path = get_runtime_base_dir() / SETTINGS_FILENAME
        self.max_proxy_attempts = DEFAULT_MAX_PROXY_ATTEMPTS
        self.thread_count_var = tk.IntVar(value=DEFAULT_THREAD_COUNT)
        self.max_proxy_attempts_var = tk.IntVar(value=DEFAULT_MAX_PROXY_ATTEMPTS)
        self.enable_proxy_precheck_var = tk.BooleanVar(value=DEFAULT_PROXY_PRECHECK)
        self.active_workers = 0
        self.enable_payment_form_var = tk.BooleanVar(value=False)
        self.enable_payment_form = False
        self.proxy_status_var = tk.StringVar(value="Proxy pool: -")

        self._build_ui()
        self._load_settings()
        self._save_settings()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.after(100, self._flush_ui_queue)

    def _build_ui(self) -> None:
        container = ttk.Frame(self.root, padding=12)
        container.pack(fill=tk.BOTH, expand=True)

        input_row = ttk.Frame(container)
        input_row.pack(fill=tk.BOTH, expand=False)

        left = ttk.LabelFrame(input_row, text="订单号 + 手机号（每行一条，支持 Tab/空格/逗号）", padding=8)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 6))
        self.orders_text = tk.Text(left, height=16, font=("Consolas", 11))
        self.orders_text.pack(fill=tk.BOTH, expand=True)
        self.orders_text.insert("1.0", DEFAULT_ORDER_PHONE_TEXT)

        right = ttk.LabelFrame(input_row, text="代理列表（每行一个，支持 ip:port 或 ip:port:user:pass）", padding=8)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(6, 0))
        self.proxies_text = tk.Text(right, height=16, font=("Consolas", 11))
        self.proxies_text.pack(fill=tk.BOTH, expand=True)
        self.proxies_text.insert("1.0", DEFAULT_PROXY_TEXT)

        btn_row = ttk.Frame(container)
        btn_row.pack(fill=tk.X, pady=10)
        self.run_btn = ttk.Button(btn_row, text="运行批量查单", command=self.on_run)
        self.run_btn.pack(side=tk.LEFT)
        self.stop_btn = ttk.Button(btn_row, text="停止", command=self.on_stop, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=8)
        self.export_btn = ttk.Button(btn_row, text="导出到桌面（xlsx）", command=self.export_manual, state=tk.DISABLED)
        self.export_btn.pack(side=tk.LEFT)
        self.payment_form_check = ttk.Checkbutton(
            btn_row,
            text="补充支付单(POST+Jobs)",
            variable=self.enable_payment_form_var,
        )
        self.payment_form_check.pack(side=tk.LEFT, padx=(10, 0))
        ttk.Label(btn_row, text="Threads").pack(side=tk.LEFT, padx=(10, 4))
        self.thread_spin = ttk.Spinbox(
            btn_row,
            from_=1,
            to=256,
            width=6,
            textvariable=self.thread_count_var,
        )
        self.thread_spin.pack(side=tk.LEFT)
        ttk.Label(btn_row, text="Retries").pack(side=tk.LEFT, padx=(10, 4))
        self.retry_spin = ttk.Spinbox(
            btn_row,
            from_=1,
            to=32,
            width=6,
            textvariable=self.max_proxy_attempts_var,
        )
        self.retry_spin.pack(side=tk.LEFT)
        self.proxy_precheck_check = ttk.Checkbutton(
            btn_row,
            text="Precheck proxies",
            variable=self.enable_proxy_precheck_var,
        )
        self.proxy_precheck_check.pack(side=tk.LEFT, padx=(10, 0))
        ttk.Label(btn_row, textvariable=self.proxy_status_var).pack(side=tk.LEFT, padx=(10, 0))

        self.progress_var = tk.StringVar(value="进度: 0/0")
        self.success_var = tk.StringVar(value="成功: 0")
        self.fail_var = tk.StringVar(value="失败: 0")
        ttk.Label(btn_row, textvariable=self.progress_var).pack(side=tk.RIGHT, padx=(12, 0))
        ttk.Label(btn_row, textvariable=self.success_var).pack(side=tk.RIGHT, padx=(12, 0))
        ttk.Label(btn_row, textvariable=self.fail_var).pack(side=tk.RIGHT, padx=(12, 0))

        failed_frame = ttk.LabelFrame(container, text="失败订单（orderId | phone | error）", padding=8)
        failed_frame.pack(fill=tk.BOTH, expand=False, pady=(0, 8))
        self.failed_text = tk.Text(failed_frame, height=6, font=("Consolas", 10), state=tk.DISABLED)
        self.failed_text.pack(fill=tk.BOTH, expand=True)

        log_frame = ttk.LabelFrame(container, text="运行日志", padding=8)
        log_frame.pack(fill=tk.BOTH, expand=True)
        self.log_text = tk.Text(log_frame, font=("Consolas", 10), state=tk.DISABLED)
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def _clamp_int(self, value: object, default: int, minimum: int, maximum: int) -> int:
        try:
            parsed = int(value)
        except Exception:
            parsed = default
        return max(minimum, min(maximum, parsed))

    def _get_proxies_text(self) -> str:
        return str(self.proxies_text.get("1.0", tk.END)).rstrip("\r\n")

    def _set_proxies_text(self, text: str) -> None:
        self.proxies_text.delete("1.0", tk.END)
        self.proxies_text.insert("1.0", text)

    def _collect_settings(self) -> dict:
        thread_count = self._clamp_int(self.thread_count_var.get(), DEFAULT_THREAD_COUNT, 1, 256)
        max_attempts = self._clamp_int(
            self.max_proxy_attempts_var.get(),
            DEFAULT_MAX_PROXY_ATTEMPTS,
            1,
            32,
        )
        self.thread_count_var.set(thread_count)
        self.max_proxy_attempts_var.set(max_attempts)
        self.max_proxy_attempts = max_attempts
        return {
            "thread_count": thread_count,
            "max_proxy_attempts": max_attempts,
            "enable_proxy_precheck": bool(self.enable_proxy_precheck_var.get()),
            "proxy_list_text": self._get_proxies_text(),
        }

    def _load_settings(self) -> None:
        settings = {}
        try:
            if self.settings_path.exists():
                settings = json.loads(self.settings_path.read_text(encoding="utf-8"))
        except Exception:
            settings = {}

        if not isinstance(settings, dict):
            settings = {}

        self.thread_count_var.set(
            self._clamp_int(settings.get("thread_count"), DEFAULT_THREAD_COUNT, 1, 256)
        )
        loaded_attempts = self._clamp_int(
            settings.get("max_proxy_attempts"),
            DEFAULT_MAX_PROXY_ATTEMPTS,
            1,
            32,
        )
        self.max_proxy_attempts = loaded_attempts
        self.max_proxy_attempts_var.set(loaded_attempts)
        self.enable_proxy_precheck_var.set(
            bool(settings.get("enable_proxy_precheck", DEFAULT_PROXY_PRECHECK))
        )
        if "proxy_list_text" in settings and isinstance(settings.get("proxy_list_text"), str):
            self._set_proxies_text(str(settings.get("proxy_list_text") or ""))

    def _save_settings(self) -> None:
        settings = self._collect_settings()
        try:
            self.settings_path.parent.mkdir(parents=True, exist_ok=True)
            temp_path = self.settings_path.with_suffix(self.settings_path.suffix + ".tmp")
            temp_path.write_text(
                json.dumps(settings, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            temp_path.replace(self.settings_path)
        except Exception as exc:
            self._push_log(f"Save settings failed: {exc}")

    def _on_close(self) -> None:
        self._save_settings()
        if self.worker and self.worker.is_alive() and not self.stop_event.is_set():
            should_exit = messagebox.askyesno(
                "Confirm Exit",
                "Task is still running. Stop the task and close the window?",
            )
            if not should_exit:
                return
            self.stop_event.set()
        self.root.destroy()

    def _dedupe_proxies(self, proxies: List[ProxyItem]) -> List[ProxyItem]:
        unique: List[ProxyItem] = []
        seen: set[str] = set()
        for proxy in proxies:
            key = str(proxy.raw or "").strip()
            if not key or key in seen:
                continue
            seen.add(key)
            unique.append(proxy)
        return unique

    def _reset_proxy_pool(self, proxies: List[ProxyItem]) -> None:
        self.proxy_pool = queue.Queue()
        for proxy in proxies:
            self.proxy_pool.put(proxy)
        self.proxy_pool_size = len(proxies)

    def _acquire_proxy(self, tried_proxy_raws: set[str]) -> Optional[ProxyItem]:
        if self.proxy_pool_size <= 0 or len(tried_proxy_raws) >= self.proxy_pool_size:
            return None

        while not self.stop_event.is_set():
            try:
                proxy = self.proxy_pool.get(timeout=0.2)
            except queue.Empty:
                continue
            if proxy.raw in tried_proxy_raws:
                self.proxy_pool.put(proxy)
                self.stop_event.wait(0.05)
                continue
            return proxy

        return None

    def _release_proxy(self, proxy: Optional[ProxyItem]) -> None:
        if proxy is not None:
            self.proxy_pool.put(proxy)

    def _sorted_snapshots(self) -> Tuple[List[dict], List[dict], List[dict]]:
        def _index_of(row: dict) -> int:
            try:
                return int(row.get("__index", 0))
            except Exception:
                return 0

        with self.results_lock:
            orders_snapshot = sorted(list(self.results_orders), key=_index_of)
            items_snapshot = sorted(list(self.results_items), key=_index_of)
            failed_snapshot = sorted(list(self.failed_rows), key=_index_of)
        return orders_snapshot, items_snapshot, failed_snapshot

    def on_run(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("Warning", "Task is already running.")
            return

        orders = parse_order_phone_lines(self.orders_text.get("1.0", tk.END))
        if not orders:
            messagebox.showerror("Error", "Order input is empty or invalid.")
            return

        try:
            requested_workers = int(self.thread_count_var.get())
        except Exception:
            messagebox.showerror("Error", "Thread count must be an integer.")
            return

        try:
            requested_attempts = int(self.max_proxy_attempts_var.get())
        except Exception:
            messagebox.showerror("Error", "Max retry count must be an integer.")
            return

        if requested_workers <= 0:
            messagebox.showerror("Error", "Thread count must be greater than 0.")
            return
        if requested_attempts <= 0:
            messagebox.showerror("Error", "Max retry count must be greater than 0.")
            return

        self.max_proxy_attempts = self._clamp_int(requested_attempts, DEFAULT_MAX_PROXY_ATTEMPTS, 1, 32)
        self.max_proxy_attempts_var.set(self.max_proxy_attempts)
        enable_proxy_precheck = bool(self.enable_proxy_precheck_var.get())
        self._save_settings()

        raw_proxies = parse_proxies(self.proxies_text.get("1.0", tk.END))
        proxies = self._dedupe_proxies(raw_proxies)
        removed_duplicates = len(raw_proxies) - len(proxies)
        self.enable_payment_form = bool(self.enable_payment_form_var.get())

        if not proxies:
            self.proxy_status_var.set("Proxy pool: 0")
            messagebox.showerror("Error", "Proxy pool is required and cannot be empty.")
            return

        actual_workers = min(requested_workers, len(proxies), len(orders))
        if actual_workers <= 0:
            messagebox.showerror("Error", "No available worker threads or proxies.")
            return

        precheck_text = "ON" if enable_proxy_precheck else "OFF"
        self.proxy_status_var.set(
            f"Proxy pool: {len(proxies)} | Threads: <= {actual_workers} | Precheck: {precheck_text}"
        )
        self.results_orders = []
        self.results_items = []
        self.failed_rows = []
        self.total = len(orders)
        self.done = 0
        self.success = 0
        self.active_workers = actual_workers
        self.last_export_path = None
        self._refresh_stats()
        self._render_failed()
        self._clear_log()

        self.run_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        self.export_btn.config(state=tk.DISABLED)
        self.stop_event.clear()

        self._push_log(
            f"Run requested: total={self.total}, requested_threads={requested_workers}, proxies={len(proxies)}, retries={self.max_proxy_attempts}, precheck={precheck_text}"
        )
        if removed_duplicates > 0:
            self._push_log(f"Removed duplicate proxies: {removed_duplicates}")
        if not enable_proxy_precheck:
            self._push_log("Proxy precheck is disabled by settings.")
        else:
            self._push_log("Proxy precheck moved to background worker; GUI remains responsive.")
        self._push_log("Strict proxy-pool mode is enabled: pool only, no local/system/direct fallback.")
        if self.enable_payment_form:
            self._push_log("Deferred payment form fetch is enabled.")

        self.worker = threading.Thread(
            target=self._run_batch_worker,
            args=(orders, proxies, requested_workers, enable_proxy_precheck),
            daemon=True,
        )
        self.worker.start()

    def _prefilter_proxies(self, orders: List[Tuple[str, str]], proxies: List[ProxyItem]) -> List[ProxyItem]:
        if not proxies or not orders:
            return proxies
        order_id, phone = orders[0]
        params = dict(COMMON_PARAMS)
        params["telephoneNumber"] = phone
        url = API_TEMPLATE.format(order_id=order_id)
        self._push_log(f"代理预检启动: {len(proxies)} 条")

        def _check(p: ProxyItem) -> bool:
            s = requests.Session()
            s.headers.update(DEFAULT_HEADERS)
            s.trust_env = False
            try:
                r = s.get(url, params=params, proxies=p.as_requests_proxy(), timeout=(4, 8))
                return r.status_code in SUMMARY_HTTP_OK
            except Exception:
                return False
            finally:
                s.close()

        keep: List[ProxyItem] = []
        max_workers = min(20, max(1, len(proxies)))
        with ThreadPoolExecutor(max_workers=max_workers) as ex:
            fut_map = {ex.submit(_check, p): p for p in proxies}
            for fut in as_completed(fut_map):
                p = fut_map[fut]
                try:
                    ok = fut.result()
                except Exception:
                    ok = False
                if ok:
                    keep.append(p)

        # 保持原始顺序
        keep_raw = {p.raw for p in keep}
        keep = [p for p in proxies if p.raw in keep_raw]
        self._push_log(f"代理预检完成: 可用 {len(keep)}/{len(proxies)}")
        return keep

    def _request_with_channel(
        self,
        session: requests.Session,
        method: str,
        url: str,
        *,
        params: Optional[dict] = None,
        json_body: Optional[dict] = None,
        proxy_cfg: Optional[Dict[str, str]] = None,
        use_env_proxy: bool = True,
        timeout: Tuple[int, int] = (8, 18),
        headers: Optional[dict] = None,
    ) -> requests.Response:
        prev = session.trust_env
        session.trust_env = use_env_proxy
        try:
            kwargs = {
                "params": params,
                "json": json_body,
                "timeout": timeout,
            }
            if headers:
                kwargs["headers"] = headers
            if proxy_cfg is not None:
                kwargs["proxies"] = proxy_cfg
            elif not use_env_proxy:
                kwargs["proxies"] = {"http": None, "https": None}
            return session.request(method, url, **kwargs)
        finally:
            session.trust_env = prev

    def _short_error_body(self, resp: requests.Response, limit: int = 220) -> str:
        try:
            body = (resp.text or "").replace("\n", " ").replace("\r", " ").strip()
        except Exception:
            body = ""
        return body[:limit]

    def _fetch_deferred_payment_form(
        self,
        session: requests.Session,
        order_id: str,
        approval_id: str,
        proxy_cfg: Optional[Dict[str, str]],
        use_env_proxy: bool,
    ) -> dict:
        out = {
            "deferredFormStatus": "",
            "deferredFormJobId": "",
            "deferredFormValidUntil": "",
            "deferredFormAction": "",
            "deferredQrCodeURL": "",
            "deferredFormError": "",
        }
        order_no = str(order_id or "").strip()
        approval = str(approval_id or "").strip()
        if not order_no or not approval:
            return out

        payload = {
            "request": {
                "approvalId": approval,
                "experienceType": "DESKTOP",
                "orderNumber": order_no,
                "returnURL": f"https://www.nike.com.cn/orders/sales/{order_no}/?payment=success",
            }
        }
        headers = dict(DEFAULT_HEADERS)
        headers["content-type"] = "application/json;charset=UTF-8"

        def _extract_job_id(node: dict) -> str:
            if not isinstance(node, dict):
                return ""
            candidates = [
                node.get("id"),
                node.get("jobId"),
                ((node.get("response") or {}).get("id") if isinstance(node.get("response"), dict) else ""),
                ((node.get("response") or {}).get("jobId") if isinstance(node.get("response"), dict) else ""),
            ]
            for raw in candidates:
                text_value = str(raw or "").strip()
                if text_value:
                    return text_value
            ref = ""
            links = node.get("links") or {}
            if isinstance(links, dict):
                self_link = links.get("self") or {}
                if isinstance(self_link, dict):
                    ref = str(self_link.get("ref") or "").strip()
            if ref:
                return ref.rstrip("/").split("/")[-1]
            return ""

        def _merge_payload(node: dict) -> None:
            if not isinstance(node, dict):
                return
            status = str(node.get("status") or "").strip()
            if status:
                out["deferredFormStatus"] = status
            maybe_job_id = _extract_job_id(node)
            if maybe_job_id and not out["deferredFormJobId"]:
                out["deferredFormJobId"] = maybe_job_id

            response_obj = node.get("response")
            if not isinstance(response_obj, dict):
                return
            qr = str(response_obj.get("qrCodeURL") or "").strip()
            if qr:
                out["deferredQrCodeURL"] = qr

            form = response_obj.get("form")
            if not isinstance(form, dict):
                return
            action = str(form.get("action") or "").strip()
            if action:
                out["deferredFormAction"] = action

            fields = form.get("fields")
            if isinstance(fields, dict):
                status_node = fields.get("status")
                if isinstance(status_node, dict):
                    valid_until = str(status_node.get("validUntil") or "").strip()
                    if valid_until:
                        out["deferredFormValidUntil"] = valid_until
                if not out["deferredFormValidUntil"]:
                    fallback_valid_until = str(fields.get("validUntil") or "").strip()
                    if fallback_valid_until:
                        out["deferredFormValidUntil"] = fallback_valid_until
            if not out["deferredFormValidUntil"]:
                response_valid_until = str(response_obj.get("validUntil") or "").strip()
                if response_valid_until:
                    out["deferredFormValidUntil"] = response_valid_until

        try:
            r = self._request_with_channel(
                session,
                "POST",
                DEFERRED_PAYMENT_POST_URL,
                json_body=payload,
                proxy_cfg=proxy_cfg,
                use_env_proxy=use_env_proxy,
                timeout=(8, 20),
                headers=headers,
            )
            if r.status_code not in (200, 202):
                out["deferredFormError"] = f"POST HTTP {r.status_code} | {self._short_error_body(r)}"
                return out
            j = r.json()
        except Exception as exc:
            out["deferredFormError"] = f"POST {type(exc).__name__}: {exc}"
            return out

        if not isinstance(j, dict):
            out["deferredFormError"] = "POST response is not json object"
            return out

        _merge_payload(j)
        job_id = str(out.get("deferredFormJobId") or "").strip()
        if not job_id:
            return out

        timed_out = True
        for _ in range(8):
            if self.stop_event.is_set():
                out["deferredFormError"] = "stopped"
                return out
            try:
                jr = self._request_with_channel(
                    session,
                    "GET",
                    DEFERRED_PAYMENT_JOB_URL.format(job_id=job_id),
                    proxy_cfg=proxy_cfg,
                    use_env_proxy=use_env_proxy,
                    timeout=(8, 20),
                )
                if jr.status_code != 200:
                    out["deferredFormError"] = f"JOB HTTP {jr.status_code} | {self._short_error_body(jr)}"
                    return out
                jj = jr.json()
            except Exception as exc:
                out["deferredFormError"] = f"JOB {type(exc).__name__}: {exc}"
                return out

            if not isinstance(jj, dict):
                out["deferredFormError"] = "JOB response is not json object"
                return out

            _merge_payload(jj)
            status_upper = str(out.get("deferredFormStatus") or "").strip().upper()
            has_required = bool(
                str(out.get("deferredQrCodeURL") or "").strip()
                and str(out.get("deferredFormAction") or "").strip()
                and str(out.get("deferredFormValidUntil") or "").strip()
            )
            if status_upper in {"FAILED", "ERROR", "CANCELLED"}:
                if not out["deferredFormError"]:
                    out["deferredFormError"] = f"JOB {status_upper}"
                return out
            if has_required and status_upper in {"COMPLETED", "SUCCESS", "DONE"}:
                timed_out = False
                break
            if status_upper in {"COMPLETED", "SUCCESS", "DONE"}:
                timed_out = False
                break
            if has_required:
                timed_out = False
                break
            self.stop_event.wait(0.4)
        else:
            timed_out = True

        if timed_out and not out["deferredFormError"]:
            out["deferredFormError"] = "JOB poll timeout"
        return out

    def on_stop(self) -> None:
        if self.worker and self.worker.is_alive():
            if not self.stop_event.is_set():
                self.stop_event.set()
                self.stop_btn.config(state=tk.DISABLED)
                self._push_log("收到停止指令，当前请求结束后退出。")

    def export_manual(self) -> None:
        orders_snapshot, items_snapshot, failed_snapshot = self._sorted_snapshots()
        if not orders_snapshot and not failed_snapshot:
            messagebox.showwarning("提示", "暂无可导出数据")
            return
        try:
            out = export_to_desktop(orders_snapshot, items_snapshot, failed_snapshot)
            self.last_export_path = out
            self._push_log(f"手动导出完成: {out}")
            messagebox.showinfo("导出成功", f"文件已导出到：\n{out}")
        except Exception as exc:
            self._push_log(f"导出失败: {exc}")
            messagebox.showerror("导出失败", str(exc))

    def _run_single_order_task(self, index: int, order_id: str, phone: str) -> dict:
        if self.stop_event.is_set():
            return {
                "index": index,
                "order_id": order_id,
                "phone": phone,
                "stopped": True,
                "ok": False,
                "row": None,
                "items": [],
                "error": "__STOP__",
                "used_proxy": "",
                "attempt": 0,
            }

        session = requests.Session()
        session.headers.update(DEFAULT_HEADERS)
        try:
            ok, row, items, err_text, used_proxy, used_attempt = self._query_single(
                session=session,
                order_id=order_id,
                phone=phone,
            )
            stopped = err_text == "__STOP__"
            return {
                "index": index,
                "order_id": order_id,
                "phone": phone,
                "stopped": stopped,
                "ok": ok,
                "row": row,
                "items": items,
                "error": err_text,
                "used_proxy": used_proxy,
                "attempt": used_attempt,
            }
        except Exception as exc:
            return {
                "index": index,
                "order_id": order_id,
                "phone": phone,
                "stopped": False,
                "ok": False,
                "row": None,
                "items": [],
                "error": f"{type(exc).__name__}: {exc}",
                "used_proxy": "",
                "attempt": 0,
            }
        finally:
            session.close()

    def _record_task_result(self, result: dict) -> None:
        if result.get("stopped"):
            return

        index = int(result.get("index", 0) or 0)
        order_id = str(result.get("order_id", "") or "")
        phone = str(result.get("phone", "") or "")
        used_proxy = str(result.get("used_proxy", "") or "")
        used_attempt = int(result.get("attempt", 0) or 0)

        if result.get("ok") and isinstance(result.get("row"), dict):
            row = dict(result["row"])
            row["__index"] = index
            items: List[dict] = []
            for item in result.get("items") or []:
                if isinstance(item, dict):
                    cloned = dict(item)
                    cloned["__index"] = index
                    items.append(cloned)
            with self.results_lock:
                self.done += 1
                self.success += 1
                self.results_orders.append(row)
                self.results_items.extend(items)
                success_count = self.success
                fail_count = len(self.failed_rows)
                done_count = self.done
            self.log_queue.put(
                ("log", f"[{done_count}/{self.total}] SUCCESS {order_id} / {phone} | proxy={used_proxy}")
            )
        else:
            fail = build_failed_row(
                order_id=order_id,
                query_phone=phone,
                error=str(result.get("error") or "UNKNOWN_ERROR"),
                proxy=used_proxy,
                attempt=used_attempt,
            )
            fail["__index"] = index
            with self.results_lock:
                self.done += 1
                self.failed_rows.append(fail)
                success_count = self.success
                fail_count = len(self.failed_rows)
                done_count = self.done
            self.log_queue.put(
                (
                    "log",
                    f"[{done_count}/{self.total}] FAIL {order_id} / {phone} | {fail['error']} | proxy={fail['proxy']}",
                )
            )

        self.log_queue.put(("progress", done_count, self.total, success_count, fail_count))

    def _run_batch_worker(
        self,
        orders: List[Tuple[str, str]],
        proxies: List[ProxyItem],
        requested_workers: int,
        enable_proxy_precheck: bool,
    ) -> None:
        stop_logged = False
        try:
            working_proxies = proxies
            precheck_text = "ON" if enable_proxy_precheck else "OFF"
            if enable_proxy_precheck:
                working_proxies = self._prefilter_proxies(orders, proxies)
                if self.stop_event.is_set():
                    self.log_queue.put(("log", "Task stopped before order queries started."))
                    self.log_queue.put(("done_error", "Task stopped before order queries started."))
                    return
                if not working_proxies:
                    self.proxy_pool_size = 0
                    self.log_queue.put(("proxy_status", "Proxy pool: 0 | Precheck: ON"))
                    self.log_queue.put(("log", "No working proxies remain after precheck."))
                    self.log_queue.put(("done_error", "No working proxies remain after precheck."))
                    return

            worker_count = min(requested_workers, len(working_proxies), len(orders))
            if worker_count <= 0:
                self.proxy_pool_size = len(working_proxies)
                self.log_queue.put(("proxy_status", f"Proxy pool: {len(working_proxies)} | Precheck: {precheck_text}"))
                self.log_queue.put(("done_error", "No available worker threads or proxies."))
                return

            self.active_workers = worker_count
            self.log_queue.put(
                (
                    "proxy_status",
                    f"Proxy pool: {len(working_proxies)} | Threads: {worker_count} | Precheck: {precheck_text}",
                )
            )
            if requested_workers > worker_count:
                self.log_queue.put(
                    (
                        "log",
                        f"Thread count reduced to {worker_count} to avoid concurrent reuse of one proxy.",
                    )
                )
            self.log_queue.put(
                (
                    "log",
                    f"Run started: total={self.total}, threads={worker_count}, proxies={len(working_proxies)}, retries={self.max_proxy_attempts}, precheck={precheck_text}",
                )
            )

            self._reset_proxy_pool(working_proxies)
            with ThreadPoolExecutor(max_workers=worker_count) as executor:
                future_map = {
                    executor.submit(self._run_single_order_task, index, order_id, phone): (index, order_id, phone)
                    for index, (order_id, phone) in enumerate(orders, start=1)
                }
                for future in as_completed(future_map):
                    try:
                        result = future.result()
                    except Exception as exc:
                        index, order_id, phone = future_map[future]
                        result = {
                            "index": index,
                            "order_id": order_id,
                            "phone": phone,
                            "stopped": False,
                            "ok": False,
                            "row": None,
                            "items": [],
                            "error": f"{type(exc).__name__}: {exc}",
                            "used_proxy": "",
                            "attempt": 0,
                        }
                    self._record_task_result(result)
                    if self.stop_event.is_set() and not stop_logged:
                        self.log_queue.put(("log", "Stop requested, waiting for active threads to finish..."))
                        stop_logged = True

            if self.stop_event.is_set():
                self.log_queue.put(("log", "Task stopped. Exporting completed records only."))

            try:
                orders_snapshot, items_snapshot, failed_snapshot = self._sorted_snapshots()
                out_path = export_to_desktop(orders_snapshot, items_snapshot, failed_snapshot)
                self.last_export_path = out_path
                self.log_queue.put(("log", f"Auto export complete: {out_path}"))
                self.log_queue.put(("done", str(out_path)))
            except Exception as exc:
                self.log_queue.put(("log", f"Auto export failed: {exc}"))
                self.log_queue.put(("done", None))
        finally:
            self.active_workers = 0

    def _query_single(
        self,
        session: requests.Session,
        order_id: str,
        phone: str,
    ) -> Tuple[bool, Optional[dict], List[dict], str, str, int]:
        last_error = "NO_PROXY_AVAILABLE"
        used_proxy_raw = ""
        final_attempt = 0
        params = dict(COMMON_PARAMS)
        params["telephoneNumber"] = phone
        url = API_TEMPLATE.format(order_id=order_id)
        tried_proxy_raws: set[str] = set()
        max_attempts = max(1, min(self.max_proxy_attempts, self.proxy_pool_size))

        for attempt in range(1, max_attempts + 1):
            if self.stop_event.is_set():
                return False, None, [], "__STOP__", used_proxy_raw, final_attempt

            proxy_item = self._acquire_proxy(tried_proxy_raws)
            if proxy_item is None:
                if self.stop_event.is_set():
                    return False, None, [], "__STOP__", used_proxy_raw, final_attempt
                break

            proxy_cfg = proxy_item.as_requests_proxy()
            used_proxy_raw = f"POOL:{proxy_item.raw}"
            final_attempt = attempt

            try:
                try:
                    r = self._request_with_channel(
                        session,
                        "GET",
                        url,
                        params=params,
                        proxy_cfg=proxy_cfg,
                        use_env_proxy=False,
                        timeout=(8, 18),
                    )
                except requests.RequestException as exc:
                    last_error = f"{type(exc).__name__}: {exc}"
                    continue

                err_body = self._short_error_body(r)
                if r.status_code != 200:
                    last_error = f"HTTP {r.status_code}" + (f" | {err_body}" if err_body else "")
                    continue

                try:
                    j = r.json()
                except Exception:
                    last_error = "INVALID_JSON"
                    continue

                if not isinstance(j, dict) or not j.get("orderId"):
                    msg = ""
                    if isinstance(j, dict):
                        msg = j.get("message") or j.get("error") or ""
                    last_error = str(msg or "ORDER_NOT_FOUND_OR_INVALID_RESPONSE")
                    continue

                row = normalize_order_row(j, phone)
                items = normalize_items(j, phone)
                row = ensure_deferred_form_fields(row)

                if self.enable_payment_form:
                    approval_id = resolve_payment_approval_id(row, j)
                    pay = self._fetch_deferred_payment_form(
                        session=session,
                        order_id=str(row.get("orderId") or order_id or "").strip(),
                        approval_id=approval_id,
                        proxy_cfg=proxy_cfg,
                        use_env_proxy=False,
                    )
                    row.update(pay)

                return True, row, items, "", used_proxy_raw, final_attempt
            finally:
                tried_proxy_raws.add(proxy_item.raw)
                self._release_proxy(proxy_item)

        return False, None, [], last_error, used_proxy_raw, final_attempt

    def _flush_ui_queue(self) -> None:
        while True:
            try:
                evt = self.log_queue.get_nowait()
            except queue.Empty:
                break

            kind = evt[0]
            if kind == "log":
                self._append_log(evt[1])
            elif kind == "progress":
                _, done, total, success, fail = evt
                self.progress_var.set(f"进度: {done}/{total}")
                self.success_var.set(f"成功: {success}")
                self.fail_var.set(f"失败: {fail}")
                self._render_failed()
            elif kind == "proxy_status":
                self.proxy_status_var.set(evt[1])
            elif kind == "done_error":
                self.run_btn.config(state=tk.NORMAL)
                self.stop_btn.config(state=tk.DISABLED)
                self.export_btn.config(state=tk.DISABLED)
                messagebox.showwarning("未开始", evt[1])
            elif kind == "done":
                self.run_btn.config(state=tk.NORMAL)
                self.stop_btn.config(state=tk.DISABLED)
                self.export_btn.config(state=tk.NORMAL)
                if evt[1]:
                    messagebox.showinfo("完成", f"查单完成：{self.success}/{self.total}\n已导出：{evt[1]}")
                else:
                    messagebox.showwarning("完成", f"查单完成：{self.success}/{self.total}\n自动导出失败，请点手动导出。")

        self.root.after(100, self._flush_ui_queue)

    def _refresh_stats(self) -> None:
        self.progress_var.set(f"进度: {self.done}/{self.total}")
        self.success_var.set(f"成功: {self.success}")
        self.fail_var.set(f"失败: {len(self.failed_rows)}")

    def _clear_log(self) -> None:
        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete("1.0", tk.END)
        self.log_text.config(state=tk.DISABLED)

    def _append_log(self, msg: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"[{ts}] {msg}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)

    def _push_log(self, msg: str) -> None:
        self.log_queue.put(("log", msg))

    def _render_failed(self) -> None:
        with self.results_lock:
            failed_snapshot = list(self.failed_rows)
        self.failed_text.config(state=tk.NORMAL)
        self.failed_text.delete("1.0", tk.END)
        for row in failed_snapshot:
            self.failed_text.insert(
                tk.END,
                f"{row.get('orderId', '')} | {row.get('queryPhone', '')} | {row.get('error', '')}\n",
            )
        self.failed_text.config(state=tk.DISABLED)


def find_desktop_dir() -> Path:
    home = Path.home()
    candidates = [
        home / "Desktop",
        home / "OneDrive" / "Desktop",
        Path(str(home).replace("\\", "/")) / "Desktop",
    ]
    for p in candidates:
        if p.exists() and p.is_dir():
            return p
    return home


def write_sheet(ws, columns: List[str], rows: List[dict]) -> None:
    ws.append(columns)
    for r in rows:
        ws.append([r.get(col, "") for col in columns])


def export_to_desktop(order_rows: List[dict], item_rows: List[dict], failed_rows: List[dict]) -> Path:
    desktop = find_desktop_dir()
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    out = desktop / f"nike_guest_orders_{ts}.xlsx"

    wb = Workbook()
    ws_orders = wb.active
    ws_orders.title = "orders"
    write_sheet(ws_orders, ORDER_COLUMNS, order_rows)

    ws_items = wb.create_sheet("items")
    write_sheet(ws_items, ITEM_COLUMNS, item_rows)

    ws_failed = wb.create_sheet("failed")
    write_sheet(ws_failed, FAILED_COLUMNS, failed_rows)

    wb.save(out)
    return out


def main() -> None:
    root = tk.Tk()
    app = NikeGuestBatchGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
